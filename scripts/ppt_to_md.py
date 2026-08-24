#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""PPTX → Markdown 转换工具：提取每页文字与表格，识别 PART 章节页，输出结构化 MD。

用法:
    python ppt_to_md.py <input.pptx> [output.md]

说明:
- 表格转为 Markdown 表格
- 识别 "PART 01/02..." 章节分隔页 → 一级标题
- 纯图片页（无文字）跳过并在文中标注
- 输出按页组织（### 第N页），便于系统切分与检索
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

PART_RE = re.compile(r"PART\s*[\n\r]*\s*(\d{1,2})")


def extract_slide(slide) -> list[str]:
    """提取一页中的文本段落与表格（Markdown 表格）。"""
    parts: list[str] = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # 组合形状（group）
                walk(sh.shapes)
                continue
            if sh.has_table:
                rows = []
                for row in sh.table.rows:
                    cells = [c.text.strip().replace("|", "\\|") for c in row.cells]
                    rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    parts.append(rows[0])
                    parts.append("|" + "|".join("---" for _ in rows[0].split("|")[1:-1]) + "|")
                    parts.extend(rows[1:])
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text.strip())
    walk(slide.shapes)
    return parts


def is_part_page(text: str) -> bool:
    """判断是否章节分隔页（含 PART NN 且文字很少）。"""
    m = PART_RE.search(text)
    return bool(m) and len(text) <= 40


def ppt_to_md(input_path: str) -> str:
    prs = Presentation(input_path)
    lines: list[str] = []
    chapter_num = 0
    for i, slide in enumerate(prs.slides, 1):
        parts = extract_slide(slide)
        text = "\n".join(p.strip() for p in parts if p.strip()).strip()
        if not text:
            lines.append(f"<!-- 第 {i} 页：纯图片，无可提取文字 -->")
            continue
        if is_part_page(text):
            m = PART_RE.search(text)
            chapter_num = int(m.group(1))
            title = re.sub(r"PART\s*[\n\r]*\s*\d{1,2}\s*", "", text).strip() or f"PART {chapter_num:02d}"
            lines.append(f"\n# PART {chapter_num:02d} {title}\n")
            continue
        lines.append(f"\n### 第 {i} 页\n\n{text}\n")
    header = "# 亚马逊售后客服培训\n\n> 来源：亚马逊售后客服培训-0115.pptx（深圳市循践量化科技有限公司）\n"
    return header + "\n".join(lines)


def main() -> None:
    if len(sys.argv) < 2:
        print("用法: python ppt_to_md.py <input.pptx> [output.md]")
        sys.exit(1)
    src = Path(sys.argv[1])
    md = ppt_to_md(str(src))
    if len(sys.argv) >= 3:
        out = Path(sys.argv[2])
        out.write_text(md, encoding="utf-8")
        print(f"已生成: {out}（{len(md)} 字符）")
    else:
        print(md)


if __name__ == "__main__":
    main()
